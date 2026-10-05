"""报告：Excel 分类清单 + Markdown 分类方案。平台无关。"""

from __future__ import annotations

import time
from collections import Counter, defaultdict

from .model import Collection, Item
from .paths import OUT
from .safety import log
from .text import fmt_duration

UNDECIDED = "待定"


def build_rows(verdicts: dict, plan: dict, review_labels: list[dict] | None,
               items: list[Item], collections: dict[str, Collection]) -> list[dict]:
    by_id = {it.id: it for it in items}
    rows: dict[str, dict] = {}

    def add(item_id: str, targets: list[str], layer: str, score, reason: str) -> None:
        item = by_id.get(item_id)
        if item is None or not targets:
            return
        rows[item_id] = {
            "id": item_id,
            "标题": item.title,
            "作者": item.creator_name,
            "时长": fmt_duration(item.duration),
            "标签": ", ".join(item.tags or []),
            "原生分类": item.category_name,
            "当前收藏夹": " / ".join(collections[c].name for c in item.containers if c in collections),
            "目标收藏夹": " + ".join(targets),
            "targets": targets,
            "判定层级": layer,
            "置信度": score if score is not None else "",
            "判定理由": reason,
            "链接": item.link,
        }

    layer_name = {"manual": "手写规则", "rules": "规则", "profile": "画像匹配", "review": "复核"}

    for a in verdicts.get("assignments") or []:
        add(a["item_id"], a.get("targets") or [], layer_name.get(a.get("layer"), a.get("layer") or ""),
            a.get("score"), a.get("reason") or "")

    for a in verdicts.get("in_place") or []:
        add(a["item_id"], a.get("targets") or [], "已就位", a.get("score"), a.get("reason") or "")

    for a in verdicts.get("unavailable") or []:
        add(a["item_id"], ["⚠ 已失效"], "已失效", "", a.get("reason") or "条目不可访问")

    for lab in review_labels or []:
        iid = lab.get("item_id") or lab.get("id")
        if not iid or iid in rows or iid not in by_id:
            continue
        ts = lab.get("targets")
        targets = [t for t in (ts if isinstance(ts, list) else [lab.get("target")]) if t]
        add(iid, targets or [UNDECIDED], "复核", lab.get("confidence"), lab.get("reason") or "")

    for it in items:
        if it.id not in rows:
            add(it.id, [UNDECIDED], "未定", "", "未参与本次分类")

    return list(rows.values())


COLUMNS = [
    ("标题", 52), ("作者", 18), ("时长", 9), ("标签", 34), ("原生分类", 14),
    ("当前收藏夹", 20), ("目标收藏夹", 26), ("判定层级", 10), ("置信度", 9),
    ("判定理由", 40), ("链接", 38),
]


def write_excel(rows: list[dict], path=None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = path or (OUT / "分类清单.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "分类清单"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append([c[0] for c in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center")

    keys = [c[0] for c in COLUMNS]
    for r in rows:
        ws.append([r.get(k, "") for k in keys])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (1, 4, 10))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("分类汇总")
    counter: Counter = Counter()
    for r in rows:
        for t in r.get("targets") or []:
            counter[t] += 1
    ws2.append(["目标收藏夹", "条目数"])
    for c in ("A1", "B1"):
        ws2[c].font = header_font
        ws2[c].fill = header_fill
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 10
    for name, n in counter.most_common():
        ws2.append([name, n])

    ws3 = wb.create_sheet("多归属条目")
    ws3.append(["标题", "目标收藏夹", "当前收藏夹"])
    for c in ("A1", "B1", "C1"):
        ws3[c].font = header_font
        ws3[c].fill = header_fill
    for col, w in (("A", 52), ("B", 34), ("C", 22)):
        ws3.column_dimensions[col].width = w
    for r in rows:
        if len(r.get("targets") or []) > 1:
            ws3.append([r["标题"], r["目标收藏夹"], r["当前收藏夹"]])

    _save(wb, path)


def _save(wb, path) -> None:
    """保存；文件被 Excel 占用时自动换名，避免整个流程失败。"""
    try:
        wb.save(path)
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{time.strftime('%H%M%S')}{path.suffix}")
        wb.save(alt)
        log.warn(f"{path.name} 正被占用，已改存为 {alt.name}")


def write_markdown(rows: list[dict], plan: dict, collections: dict[str, Collection],
                   path=None) -> None:
    path = path or (OUT / "分类方案.md")
    grouped: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        for t in r.get("targets") or []:
            grouped[t].append(r)

    lines = ["# 收藏夹分类方案", "",
             f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}", "",
             "## 总览", "",
             f"- 现有收藏夹：**{len(collections)}** 个",
             f"- 条目总数：**{len(rows)}**"]
    s = plan.get("summary") or {}
    if s:
        lines += [
            f"- 需要处理：**{s.get('items_to_touch', 0)}** 个，共 {s.get('total_add_ops', 0)} 次收藏操作",
            f"- 归入多个收藏夹：**{s.get('multi_target_items', 0)}** 个",
            f"- 移出收件箱：**{s.get('items_leaving_inbox', 0)}** 个",
            f"- 已失效不动：**{s.get('unavailable', 0)}** 个",
        ]
    lines += ["", "> 采用「追加收藏」：条目被加入目标收藏夹，原有收藏夹保持不变，",
              "> 只有收件箱（默认收藏夹）会在归类后被清空。", ""]

    lines += ["## 分类分布", "", "| 目标收藏夹 | 条目数 |", "|---|---:|"]
    for name, its in sorted(grouped.items(), key=lambda kv: -len(kv[1])):
        lines.append(f"| {name} | {len(its)} |")
    lines.append("")

    if plan.get("collections_to_create"):
        lines += ["## 需要新建的收藏夹", ""]
        lines += [f"- {n}" for n in plan["collections_to_create"]]
        lines.append("")

    lines += ["## 明细（按目标分类）", ""]
    for name, its in sorted(grouped.items(), key=lambda kv: -len(kv[1])):
        lines.append(f"### {name}（{len(its)}）")
        lines.append("")
        lines.append("| 标题 | 作者 | 时长 | 当前收藏夹 | 全部目标 | 层级 |")
        lines.append("|---|---|---|---|---|---|")
        for r in its[:200]:
            title = (r["标题"] or "").replace("|", "丨")
            lines.append(
                f"| [{title}]({r['链接']}) | {r['作者']} | {r['时长']} | "
                f"{r['当前收藏夹']} | {r['目标收藏夹']} | {r['判定层级']} |"
            )
        if len(its) > 200:
            lines.append(f"| …另有 {len(its) - 200} 条，见 Excel | | | | | |")
        lines.append("")

    path.write_text("\n".join(lines), encoding="utf-8")


def report(verdicts: dict, plan: dict, review_labels, items, collections) -> None:
    log.step("生成报告…")
    rows = build_rows(verdicts, plan, review_labels, items, collections)
    write_excel(rows)
    write_markdown(rows, plan, collections)
    log.ok(f"Excel：{OUT / '分类清单.xlsx'}")
    log.ok(f"Markdown：{OUT / '分类方案.md'}")

    counter: Counter = Counter()
    for r in rows:
        for t in r.get("targets") or []:
            counter[t] += 1
    multi = sum(1 for r in rows if len(r.get("targets") or []) > 1)
    log.info("")
    log.info(f"  多归属条目：{multi} 个")
    for name, n in counter.most_common(30):
        log.info(f"    {n:>5}  {name}")
